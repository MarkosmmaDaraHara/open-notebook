from __future__ import annotations

import html
import json
import os
import re
import threading
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable
from uuid import uuid4

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pydantic import BaseModel, Field

DEFAULT_DOCUMENT_HTML = """
<h1>Project proposal</h1>
<p class="subtitle">A practical plan for the next release</p>
<h2>1. Executive summary</h2>
<p>This proposal outlines the plan, resources, and timeline for delivering the next release. The focus is on solving the most important customer problems while improving reliability and performance. We will ship valuable features, reduce risk, and establish a strong foundation for future iterations.</p>
<blockquote>Our approach balances short-term impact with long-term sustainability. We will prioritize high-value work, collaborate closely across teams, and maintain a clear feedback loop with users. Success will be measured by adoption, performance improvements, and customer satisfaction.</blockquote>
<h2>2. Objectives</h2>
<p>These objectives guide our decisions and help us measure progress throughout the release cycle.</p>
<table>
  <thead><tr><th>Objective</th><th>Description</th><th>Metric</th></tr></thead>
  <tbody>
    <tr><td>Deliver value</td><td>Ship features that solve key user problems</td><td>Adoption rate</td></tr>
    <tr><td>Improve quality</td><td>Enhance reliability and performance</td><td>Error rate, uptime</td></tr>
    <tr><td>Increase efficiency</td><td>Streamline workflows and reduce cycle time</td><td>Cycle time, throughput</td></tr>
  </tbody>
</table>
<p>By focusing on these objectives, we will create meaningful outcomes for users and the business.</p>
""".strip()


class DocumentNotFoundError(LookupError):
    """Raised when a document identifier does not exist."""


class DocumentVersionConflictError(RuntimeError):
    """Raised when an optimistic update uses an out-of-date document version."""


class DocumentRecord(BaseModel):
    id: str
    title: str
    file_type: str = "docx"
    html: str
    version: int = 1
    created_at: datetime
    updated_at: datetime
    word_count: int = Field(default=0, ge=0)


def html_to_plain_text(content: str) -> str:
    soup = BeautifulSoup(content, "html.parser")
    return "\n".join(
        line.strip() for line in soup.get_text("\n").splitlines() if line.strip()
    )


def _word_count(content: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", html_to_plain_text(content), re.UNICODE))


def _sanitize_title(title: str) -> str:
    value = " ".join(title.strip().split())
    value = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "", value)
    value = " ".join(value.split())
    return value[:160] or "Untitled document"


def _sanitize_html(content: str) -> str:
    """Remove executable markup while retaining office-style formatting."""

    soup = BeautifulSoup(content or "<p></p>", "html.parser")
    for element in soup.find_all(["script", "iframe", "object", "embed", "form"]):
        element.decompose()

    safe_url_schemes = ("http://", "https://", "mailto:", "data:image/")
    for element in soup.find_all(True):
        for attribute in list(element.attrs):
            lowered = attribute.lower()
            if lowered.startswith("on") or lowered in {"srcdoc", "formaction"}:
                del element.attrs[attribute]
        for attribute in ("href", "src"):
            value = element.attrs.get(attribute)
            if (
                isinstance(value, str)
                and value
                and not value.startswith(safe_url_schemes)
            ):
                if not value.startswith(("/", "#")):
                    del element.attrs[attribute]

    return str(soup)


class LocalDocumentStore:
    """A durable local adapter designed to be swapped for S3/object storage later."""

    def __init__(self, root: str | Path | None = None):
        configured_root = root or os.environ.get(
            "IQ_DOCUMENTS_DIR", "data/iq-documents"
        )
        self.root = Path(configured_root).expanduser().resolve()
        self.metadata_dir = self.root / "metadata"
        self.file_dir = self.root / "files"
        self.metadata_dir.mkdir(parents=True, exist_ok=True)
        self.file_dir.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()

    @staticmethod
    def _validate_id(document_id: str) -> str:
        if not re.fullmatch(r"[a-f0-9-]{20,64}", document_id):
            raise DocumentNotFoundError(document_id)
        return document_id

    def _metadata_path(self, document_id: str) -> Path:
        return self.metadata_dir / f"{self._validate_id(document_id)}.json"

    def docx_path(self, document_id: str) -> Path:
        self._validate_id(document_id)
        return self.file_dir / f"{document_id}.docx"

    def _write_record(self, record: DocumentRecord) -> None:
        target = self._metadata_path(record.id)
        temporary = target.with_suffix(".json.tmp")
        temporary.write_text(record.model_dump_json(indent=2), encoding="utf-8")
        temporary.replace(target)

    def _read_record(self, document_id: str) -> DocumentRecord:
        path = self._metadata_path(document_id)
        if not path.exists():
            raise DocumentNotFoundError(document_id)
        return DocumentRecord.model_validate_json(path.read_text(encoding="utf-8"))

    def list_documents(self, query: str | None = None) -> list[DocumentRecord]:
        query_value = (query or "").strip().casefold()
        records: list[DocumentRecord] = []
        with self._lock:
            for path in self.metadata_dir.glob("*.json"):
                try:
                    record = DocumentRecord.model_validate_json(
                        path.read_text(encoding="utf-8")
                    )
                except Exception:
                    continue
                if query_value:
                    haystack = (
                        f"{record.title}\n{html_to_plain_text(record.html)}".casefold()
                    )
                    if query_value not in haystack:
                        continue
                records.append(record)
        return sorted(records, key=lambda item: item.updated_at, reverse=True)

    def get_document(self, document_id: str) -> DocumentRecord:
        with self._lock:
            return self._read_record(document_id)

    def create_document(
        self,
        title: str = "Untitled document",
        initial_html: str | None = None,
        initial_text: str | None = None,
    ) -> DocumentRecord:
        now = datetime.now(UTC)
        document_id = str(uuid4())
        if initial_html is not None:
            content = _sanitize_html(initial_html)
        elif initial_text is not None:
            paragraphs = [
                f"<p>{html.escape(line)}</p>"
                for line in initial_text.splitlines()
                if line
            ]
            content = "".join(paragraphs) or "<p></p>"
        else:
            content = DEFAULT_DOCUMENT_HTML

        record = DocumentRecord(
            id=document_id,
            title=_sanitize_title(title),
            html=content,
            version=1,
            created_at=now,
            updated_at=now,
            word_count=_word_count(content),
        )
        with self._lock:
            self._write_record(record)
            self._write_docx(record)
        return record

    def update_document(
        self,
        document_id: str,
        *,
        title: str | None = None,
        content_html: str | None = None,
        expected_version: int | None = None,
    ) -> DocumentRecord:
        with self._lock:
            record = self._read_record(document_id)
            if expected_version is not None and record.version != expected_version:
                raise DocumentVersionConflictError(
                    f"Expected version {expected_version}, found {record.version}"
                )

            next_title = _sanitize_title(title) if title is not None else record.title
            next_html = (
                _sanitize_html(content_html)
                if content_html is not None
                else record.html
            )
            if next_title == record.title and next_html == record.html:
                return record

            record.title = next_title
            record.html = next_html
            record.word_count = _word_count(next_html)
            record.version += 1
            record.updated_at = datetime.now(UTC)
            self._write_record(record)
            self._write_docx(record)
            return record

    def delete_document(self, document_id: str) -> None:
        with self._lock:
            metadata_path = self._metadata_path(document_id)
            if not metadata_path.exists():
                raise DocumentNotFoundError(document_id)
            metadata_path.unlink()
            file_path = self.docx_path(document_id)
            if file_path.exists():
                file_path.unlink()

    def append_text(self, document_id: str, text: str) -> DocumentRecord:
        record = self.get_document(document_id)
        soup = BeautifulSoup(record.html, "html.parser")
        for line in text.splitlines() or [text]:
            paragraph = soup.new_tag("p")
            paragraph.string = line
            soup.append(paragraph)
        return self.update_document(
            document_id, content_html=str(soup), expected_version=record.version
        )

    def replace_text(
        self,
        document_id: str,
        find: str,
        replacement: str,
        *,
        replace_all: bool = False,
        expected_version: int | None = None,
    ) -> tuple[DocumentRecord, int]:
        if not find:
            raise ValueError("find must not be empty")

        with self._lock:
            record = self._read_record(document_id)
            if expected_version is not None and record.version != expected_version:
                raise DocumentVersionConflictError(
                    f"Expected version {expected_version}, found {record.version}"
                )

            soup = BeautifulSoup(record.html, "html.parser")
            replacement_count = 0
            for node in list(soup.find_all(string=True)):
                if not isinstance(node, NavigableString) or find not in str(node):
                    continue
                limit = -1 if replace_all else 1
                replaced = str(node).replace(find, replacement, limit)
                delta = str(node).count(find) if replace_all else 1
                node.replace_with(replaced)
                replacement_count += delta
                if not replace_all:
                    break

            if replacement_count == 0:
                return record, 0
            updated = self.update_document(
                document_id,
                content_html=str(soup),
                expected_version=record.version,
            )
            return updated, replacement_count

    def save_docx_bytes(self, document_id: str, payload: bytes) -> DocumentRecord:
        if not payload:
            raise ValueError("The saved DOCX payload is empty")
        with self._lock:
            record = self._read_record(document_id)
            target = self.docx_path(document_id)
            temporary = target.with_suffix(".docx.tmp")
            temporary.write_bytes(payload)
            temporary.replace(target)
            record.version += 1
            record.updated_at = datetime.now(UTC)
            record.html = self._docx_to_html(target)
            record.word_count = _word_count(record.html)
            self._write_record(record)
            return record

    def _write_docx(self, record: DocumentRecord) -> None:
        doc = DocxDocument()
        section = doc.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        styles = doc.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(11)

        soup = BeautifulSoup(record.html, "html.parser")
        for element in self._top_level_elements(soup):
            if element.name in {"h1", "h2", "h3"}:
                level = int(element.name[1])
                paragraph = doc.add_heading(level=min(level, 3))
                self._append_inline_runs(paragraph, element)
            elif element.name == "p":
                paragraph = doc.add_paragraph()
                if "subtitle" in element.get("class", []):
                    paragraph.style = doc.styles["Subtitle"]
                if element.get("style", "").find("text-align: center") >= 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._append_inline_runs(paragraph, element)
            elif element.name in {"ul", "ol"}:
                style = "List Bullet" if element.name == "ul" else "List Number"
                for item in element.find_all("li", recursive=False):
                    paragraph = doc.add_paragraph(style=style)
                    self._append_inline_runs(paragraph, item)
            elif element.name == "blockquote":
                paragraph = doc.add_paragraph(style="Quote")
                self._append_inline_runs(paragraph, element)
            elif element.name == "table":
                rows = element.find_all("tr")
                column_count = max(
                    (len(row.find_all(["th", "td"])) for row in rows), default=1
                )
                table = doc.add_table(rows=0, cols=column_count)
                table.style = "Table Grid"
                for row_element in rows:
                    cells = row_element.find_all(["th", "td"])
                    row = table.add_row()
                    for index, cell_element in enumerate(cells[:column_count]):
                        cell = row.cells[index]
                        cell.text = cell_element.get_text(" ", strip=True)
                        if cell_element.name == "th":
                            for run in cell.paragraphs[0].runs:
                                run.bold = True

        target = self.docx_path(record.id)
        temporary = target.with_suffix(".docx.tmp")
        doc.save(temporary)
        temporary.replace(target)

    @staticmethod
    def _top_level_elements(soup: BeautifulSoup) -> Iterable[Tag]:
        root = soup.body if soup.body else soup
        return (child for child in root.children if isinstance(child, Tag))

    @classmethod
    def _append_inline_runs(cls, paragraph, element: Tag) -> None:
        def walk(
            node: Tag | NavigableString, bold: bool = False, italic: bool = False
        ) -> None:
            if isinstance(node, NavigableString):
                if str(node):
                    run = paragraph.add_run(str(node))
                    run.bold = bold
                    run.italic = italic
                return
            next_bold = bold or node.name in {"strong", "b"}
            next_italic = italic or node.name in {"em", "i"}
            if node.name == "br":
                paragraph.add_run().add_break()
                return
            for child in node.children:
                walk(child, next_bold, next_italic)

        walk(element)

    @staticmethod
    def _docx_to_html(path: Path) -> str:
        doc = DocxDocument(path)
        parts: list[str] = []
        for paragraph in doc.paragraphs:
            text = html.escape(paragraph.text)
            style_name = paragraph.style.name.casefold() if paragraph.style else ""
            if style_name.startswith("heading 1"):
                parts.append(f"<h1>{text}</h1>")
            elif style_name.startswith("heading 2"):
                parts.append(f"<h2>{text}</h2>")
            elif style_name.startswith("heading 3"):
                parts.append(f"<h3>{text}</h3>")
            elif text:
                parts.append(f"<p>{text}</p>")
        for table in doc.tables:
            parts.append("<table><tbody>")
            for row in table.rows:
                parts.append("<tr>")
                parts.extend(f"<td>{html.escape(cell.text)}</td>" for cell in row.cells)
                parts.append("</tr>")
            parts.append("</tbody></table>")
        return "".join(parts) or "<p></p>"
